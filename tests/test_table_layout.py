from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Pt

from sfu_converter.infrastructure.docx_renderer import DocxRenderer
from sfu_converter.parser import V2Parser
from sfu_converter.registry import get_profile


def _parse_v2(source: str):
    result = V2Parser().parse(source)
    assert result.diagnostics == []
    return result.document


def _table_border_values(table):
    borders = table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders is not None
    return {element.tag.rsplit("}", 1)[-1]: element.get(qn("w:val")) for element in borders}


def _cell_border_values(cell):
    borders = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
    assert borders is not None
    return {element.tag.rsplit("}", 1)[-1]: element.get(qn("w:val")) for element in borders}


def test_renderer_outputs_compliant_table_unit_header_borders_and_notes(tmp_path):
    document = _parse_v2(
        "\n".join(
            [
                '[TABLE caption="Параметры" unit="МПа" column_units="-,Гц,°С" header_rows=2]',
                "| Параметр | Частота | Температура |",
                "|----------|---------|-------------|",
                "| Обозначение | f | T |",
                "| Давление | 50 | 20 |",
                '[TABLE_NOTE marker="*" text="Значение приведено при нормальных условиях"]',
                "[TABLE_END]",
            ]
        )
    )
    output = tmp_path / "table_layout.docx"

    DocxRenderer(base_dir=tmp_path).render_to_file(document, get_profile("common"), str(output))

    doc = DocxDocument(str(output))
    unit = next(paragraph for paragraph in doc.paragraphs if paragraph.text == ", МПа")
    assert unit.style.name == "SFUTableUnit"
    assert unit.paragraph_format.alignment.name == "RIGHT"
    assert unit.runs[0].font.size == Pt(12)
    # The unit label is part of the caption+table block and must not be orphaned.
    assert unit.paragraph_format.keep_with_next is True

    table = doc.tables[0]
    assert table.style.name == "SFUTable"
    assert "Table Grid" not in table.style.name
    assert _table_border_values(table) == {
        "top": "single",
        "left": "single",
        "bottom": "single",
        "right": "single",
        "insideH": "single",
        "insideV": "single",
    }
    assert 'w:val="double"' in table.rows[1]._tr.xml
    assert "tblHeader" in table.rows[0]._tr.xml
    assert "tblHeader" in table.rows[1]._tr.xml
    assert table.rows[0].cells[1].text == "Частота, Гц"
    assert table.rows[0].cells[2].text == "Температура, °С"

    note_cell = table.rows[-1].cells[0]
    assert "Значение приведено при нормальных условиях" in note_cell.text
    assert note_cell.paragraphs[0].runs[0].font.superscript is True


def test_renderer_outputs_cell_borders_on_every_side(tmp_path):
    document = _parse_v2(
        "\n".join(
            [
                '[TABLE caption="Параметры" header_rows=1]',
                "| Название | Значение |",
                "|----------|----------|",
                "| Давление | 10 |",
                "[TABLE_END]",
            ]
        )
    )
    output = tmp_path / "table_cell_borders.docx"

    DocxRenderer(base_dir=tmp_path).render_to_file(document, get_profile("common"), str(output))

    table = DocxDocument(str(output)).tables[0]
    header = _cell_border_values(table.rows[0].cells[0])
    body = _cell_border_values(table.rows[1].cells[0])
    assert header["top"] == "single"
    assert header["left"] == "single"
    assert header["right"] == "single"
    assert header["bottom"] == "double"
    assert body == {
        "top": "single",
        "left": "single",
        "bottom": "single",
        "right": "single",
    }


def test_renderer_outputs_continuation_label_and_numbering_row(tmp_path):
    document = _parse_v2(
        "\n".join(
            [
                '[TABLE caption="Параметры" number=3 continuation=final]',
                "| Параметр | Значение |",
                "|----------|----------|",
                "| Давление | 10 |",
                "[TABLE_END]",
            ]
        )
    )
    output = tmp_path / "table_continuation.docx"

    DocxRenderer(base_dir=tmp_path).render_to_file(document, get_profile("common"), str(output))

    doc = DocxDocument(str(output))
    assert "Окончание таблицы 3" in [paragraph.text for paragraph in doc.paragraphs]
    assert [cell.text for cell in doc.tables[0].rows[0].cells] == ["1", "2"]
    assert not any(paragraph.text == "Таблица 3 — Параметры" for paragraph in doc.paragraphs)

    label = next(p for p in doc.paragraphs if p.text == "Окончание таблицы 3")
    assert label.style.name == "SFUTableCaption"
    # The continuation caption must stay on the same page as its table part.
    assert label.paragraph_format.keep_with_next is True


def test_main_caption_keeps_with_following_table(tmp_path):
    document = _parse_v2(
        "\n".join(
            [
                '[TABLE caption="Параметры" number=1]',
                "| Параметр | Значение |",
                "| Давление | 10 |",
                "[TABLE_END]",
            ]
        )
    )
    output = tmp_path / "table_caption_keep.docx"

    DocxRenderer(base_dir=tmp_path).render_to_file(document, get_profile("common"), str(output))

    doc = DocxDocument(str(output))
    caption = next(p for p in doc.paragraphs if p.text == "Таблица 1 — Параметры")
    assert caption.paragraph_format.keep_with_next is True


def _render(tmp_path, source: str, name: str = "table.docx"):
    document = _parse_v2(source)
    output = tmp_path / name
    DocxRenderer(base_dir=tmp_path).render_to_file(document, get_profile("common"), str(output))
    return DocxDocument(str(output))


def _grid_widths(table) -> list[int]:
    grid = table._tbl.find(qn("w:tblGrid"))
    return [int(column.get(qn("w:w"))) for column in grid.findall(qn("w:gridCol"))]


def _text_area_width(doc) -> int:
    section = doc.sections[-1]
    return int((section.page_width - section.left_margin - section.right_margin) / 635)


def _cell_runs(cell):
    return [(run.text, bool(run.bold)) for paragraph in cell.paragraphs for run in paragraph.runs]


def test_column_widths_grow_with_content_and_fill_the_text_area(tmp_path):
    doc = _render(
        tmp_path,
        "\n".join(
            [
                '[TABLE caption="Параметры"]',
                "| № | Наименование параметра | Значение |",
                "| 1 | Давление в камере сгорания на номинальном режиме работы | 10 |",
                "| 2 | Температура | 20 |",
                "[TABLE_END]",
            ]
        ),
    )

    widths = _grid_widths(doc.tables[0])
    assert sum(widths) == _text_area_width(doc)
    # The descriptive column carries the text, the counter stays narrow.
    assert widths[1] > widths[0] * 3
    assert widths[1] > widths[2]
    assert doc.tables[0]._tbl.tblPr.find(qn("w:tblW")).get(qn("w:type")) == "dxa"


def test_short_cells_keep_balanced_columns(tmp_path):
    doc = _render(
        tmp_path,
        "\n".join(
            [
                "[TABLE]",
                "| a | bb | cc |",
                "| 1 | 22 | 33 |",
                "[TABLE_END]",
            ]
        ),
    )

    widths = _grid_widths(doc.tables[0])
    assert sum(widths) == _text_area_width(doc)
    assert max(widths) / min(widths) < 2


def test_bold_markup_inside_cells_becomes_a_bold_run(tmp_path):
    doc = _render(
        tmp_path,
        "\n".join(
            [
                "[TABLE]",
                "| Параметр | Комментарий |",
                "| Давление | **критично** при нагреве |",
                "| Температура | обычное значение |",
                "[TABLE_END]",
            ]
        ),
    )

    table = doc.tables[0]
    assert _cell_runs(table.rows[1].cells[1]) == [("критично", True), (" при нагреве", False)]
    assert _cell_runs(table.rows[2].cells[1]) == [("обычное значение", False)]
    assert all("*" not in cell.text for row in table.rows for cell in row.cells)


def test_header_stays_bold_and_italic_letter_rule_is_kept(tmp_path):
    doc = _render(
        tmp_path,
        "\n".join(
            [
                "[TABLE]",
                "| Параметр | **Обозначение** | Значение |",
                "| Давление | p | 10 |",
                "| Нагрев | *T* | 20 |",
                "[TABLE_END]",
            ]
        ),
    )

    table = doc.tables[0]
    assert _cell_runs(table.rows[1].cells[1]) == [("p", False)]
    assert table.rows[1].cells[1].paragraphs[0].runs[0].italic is True
    assert table.rows[2].cells[1].paragraphs[0].runs[0].italic is True
    header_cell = table.rows[0].cells[1]
    assert _cell_runs(header_cell) == [("Обозначение", True)]


def test_note_row_spans_the_full_table_width(tmp_path):
    doc = _render(
        tmp_path,
        "\n".join(
            [
                "[TABLE]",
                "| Параметр | Значение |",
                "| Давление | 10 |",
                '[TABLE_NOTE marker="*" text="Данные приведены на конец периода."]',
                "[TABLE_END]",
            ]
        ),
    )

    table = doc.tables[0]
    widths = _grid_widths(table)
    note_cell = table.rows[-1]._tr.tc_lst[0]
    assert note_cell.grid_span == len(widths)
    assert note_cell.width.twips == sum(widths)


def test_autofit_false_keeps_equal_columns(tmp_path):
    doc = _render(
        tmp_path,
        "\n".join(
            [
                '[TABLE caption="Параметры" autofit=false]',
                "| № | Наименование параметра | Значение |",
                "| 1 | Давление в камере сгорания на номинальном режиме | 10 |",
                "[TABLE_END]",
            ]
        ),
    )

    widths = _grid_widths(doc.tables[0])
    assert len(set(widths)) == 1
    assert doc.tables[0]._tbl.tblPr.find(qn("w:tblW")).get(qn("w:type")) == "auto"
