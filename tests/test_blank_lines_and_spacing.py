"""Blank-line separators and single-spacing rules for structural elements."""

from __future__ import annotations

from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from PIL import Image

from sfu_converter.domain.ast_nodes import (
    BibliographyEntryNode,
    Document,
    FigureNode,
    FormulaNode,
    HeadingLevel,
    HeadingNode,
    ListItemNode,
    ListNode,
    ListType,
    ParagraphNode,
    StructuralSectionNode,
    StructuralSectionType,
    TableCell,
    TableNode,
    TableRow,
    TextRun,
)
from sfu_converter.infrastructure.docx_renderer import DocxRenderer
from sfu_converter.registry import get_profile


def _render(tmp_path, blocks, name="spacing.docx"):
    path = tmp_path / name
    DocxRenderer(base_dir=tmp_path).render_to_file(
        Document(blocks=tuple(blocks)),
        get_profile("common"),
        str(path),
    )
    return DocxDocument(str(path))


def _has_page_break(paragraph: Paragraph) -> bool:
    for br in paragraph._p.iter(qn("w:br")):
        if br.get(qn("w:type")) == "page":
            return True
    p_pr = paragraph._p.find(qn("w:pPr"))
    return p_pr is not None and p_pr.find(qn("w:sectPr")) is not None


def _blank_paragraphs(doc):
    return [
        paragraph
        for paragraph in doc.paragraphs
        if not paragraph.text.strip() and not _has_page_break(paragraph)
    ]


def test_blank_line_after_heading_but_not_between_paragraphs(tmp_path):
    doc = _render(
        tmp_path,
        (
            HeadingNode(level=HeadingLevel.H1, text="Глава", number="auto"),
            ParagraphNode(runs=(TextRun("Первый абзац."),)),
            ParagraphNode(runs=(TextRun("Второй абзац."),)),
        ),
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts == ["1 Глава", "", "Первый абзац.", "Второй абзац."]

    heading = doc.paragraphs[0]
    assert heading.paragraph_format.space_before == Pt(0)
    assert heading.paragraph_format.space_after == Pt(0)

    for blank in _blank_paragraphs(doc):
        pf = blank.paragraph_format
        assert pf.line_spacing == 1.0
        assert pf.space_before == Pt(0)
        assert pf.space_after == Pt(0)
        assert abs(pf.first_line_indent - Cm(0)) < 1000
        assert blank.runs and blank.runs[0].font.size == Pt(14)


def test_no_blank_line_between_paragraphs_and_lists(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("До списка."),)),
            ListNode(
                list_type=ListType.NUMBERED,
                items=(ListItemNode("Пункт один."), ListItemNode("Пункт два.")),
            ),
            ParagraphNode(runs=(TextRun("После списка."),)),
        ),
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts == [
        "До списка.",
        "Пункт один.",
        "Пункт два.",
        "После списка.",
    ]


def test_blank_line_after_heading_before_list(tmp_path):
    doc = _render(
        tmp_path,
        (
            HeadingNode(level=HeadingLevel.H1, text="Глава", number="auto"),
            ListNode(
                list_type=ListType.NUMBERED,
                items=(ListItemNode("Пункт один."), ListItemNode("Пункт два.")),
            ),
        ),
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts == ["1 Глава", "", "Пункт один.", "Пункт два."]


def test_no_blank_line_before_heading_at_page_start(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Текст до разрыва."),)),
            HeadingNode(level=HeadingLevel.H1, text="Новая глава", number="auto"),
            ParagraphNode(runs=(TextRun("Текст после заголовка."),)),
        ),
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    heading_index = texts.index("1 Новая глава")
    assert texts[heading_index - 1] == "Текст до разрыва."
    assert doc.paragraphs[heading_index].paragraph_format.page_break_before is True
    assert texts[heading_index + 1] == ""
    assert texts[heading_index + 2] == "Текст после заголовка."


def test_figure_and_table_captions_use_single_spacing(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("До рисунка."),)),
            FigureNode(src=None, caption="Схема", id="fig:one"),
            TableNode(
                caption="Данные",
                rows=(
                    TableRow(cells=(TableCell("A"), TableCell("B"))),
                    TableRow(cells=(TableCell("1"), TableCell("2"))),
                ),
            ),
        ),
    )

    figure_caption = next(p for p in doc.paragraphs if p.text.startswith("Рисунок 1"))
    table_caption = next(p for p in doc.paragraphs if p.text.startswith("Таблица 1"))
    for caption in (figure_caption, table_caption):
        assert caption.paragraph_format.line_spacing == 1.0
        assert caption.paragraph_format.space_before == Pt(0)
        assert caption.paragraph_format.space_after == Pt(0)

    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts[texts.index("До рисунка.") + 1] == ""
    assert texts[texts.index(figure_caption.text) + 1] == ""
    assert texts[texts.index(table_caption.text) - 1] == ""


def test_formula_block_has_blank_lines_around_it_and_no_explicit_spacing(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("До формулы."),)),
            FormulaNode(content="x = y", id="f1"),
            ParagraphNode(runs=(TextRun("После формулы."),)),
        ),
    )

    formula = next(p for p in doc.paragraphs if p.text.endswith("(1)"))
    assert formula.paragraph_format.space_before == Pt(0)
    assert formula.paragraph_format.space_after == Pt(0)

    texts = [paragraph.text for paragraph in doc.paragraphs]
    assert texts[texts.index(formula.text) - 1] == ""
    assert texts[texts.index(formula.text) + 1] == ""


def test_sources_heading_and_entries_are_separated_only_after_heading(tmp_path):
    doc = _render(
        tmp_path,
        (
            StructuralSectionNode(
                section_type=StructuralSectionType.SOURCES,
                title="СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
            ),
            BibliographyEntryNode(number=1, text="Первый источник."),
            BibliographyEntryNode(number=2, text="Второй источник."),
        ),
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    heading_index = texts.index("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")
    assert texts[heading_index + 1] == ""
    assert texts[heading_index + 2] == "Первый источник."
    assert texts[heading_index + 3] == "Второй источник."


def _marker_texts_for_paragraph(doc, paragraph):
    num_pr = paragraph._p.pPr.numPr
    num_id = num_pr.numId.val
    numbering = doc.part.numbering_part.element
    abstract_id = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(num_id):
            abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
            break
    assert abstract_id is not None
    for abstract in numbering.findall(qn("w:abstractNum")):
        if abstract.get(qn("w:abstractNumId")) == abstract_id:
            return [element.get(qn("w:val")) for element in abstract.iter(qn("w:lvlText"))]
    raise AssertionError("abstract numbering definition not found")


def test_numbered_list_in_sources_section_uses_dot_markers(tmp_path):
    doc = _render(
        tmp_path,
        (
            StructuralSectionNode(
                section_type=StructuralSectionType.SOURCES,
                title="СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
            ),
            ListNode(
                list_type=ListType.NUMBERED,
                items=(ListItemNode("Первый источник."), ListItemNode("Второй источник.")),
            ),
        ),
        name="sources_list.docx",
    )

    item = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "Первый источник.")
    assert _marker_texts_for_paragraph(doc, item)[0] == "%1."


def test_regular_numbered_list_keeps_parenthesis_marker(tmp_path):
    doc = _render(
        tmp_path,
        (
            ListNode(
                list_type=ListType.NUMBERED,
                items=(ListItemNode("Первый пункт."), ListItemNode("Второй пункт.")),
            ),
        ),
        name="regular_list.docx",
    )

    item = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "Первый пункт.")
    assert _marker_texts_for_paragraph(doc, item)[0] == "%1)"


def test_bullet_list_uses_en_dash_marker(tmp_path):
    doc = _render(
        tmp_path,
        (
            ListNode(
                list_type=ListType.BULLET,
                items=(ListItemNode("Первый пункт."), ListItemNode("Второй пункт.")),
            ),
        ),
        name="bullet_list.docx",
    )

    item = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "Первый пункт.")
    assert _marker_texts_for_paragraph(doc, item) == ["\u2013", "\u2013", "\u2013"]


def _paragraph_with_drawing(doc):
    return next(
        paragraph for paragraph in doc.paragraphs if paragraph._p.find(".//" + qn("w:drawing")) is not None
    )


def test_heading_keeps_with_next_through_blank_line(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Текст до."),)),
            HeadingNode(level=HeadingLevel.H2, text="Раздел", number="auto"),
            ParagraphNode(runs=(TextRun("Текст после."),)),
        ),
        name="keep_heading.docx",
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    heading = doc.paragraphs[texts.index("1.1 Раздел")]
    assert heading.paragraph_format.keep_with_next is True
    assert heading.paragraph_format.keep_together is True

    blank_after = doc.paragraphs[texts.index("1.1 Раздел") + 1]
    assert blank_after.text == ""
    assert blank_after.paragraph_format.keep_with_next is True

    blank_before = doc.paragraphs[texts.index("1.1 Раздел") - 1]
    assert blank_before.text == ""
    assert blank_before.paragraph_format.keep_with_next is not True


def test_sources_heading_keeps_with_next_through_blank_line(tmp_path):
    doc = _render(
        tmp_path,
        (
            StructuralSectionNode(
                section_type=StructuralSectionType.SOURCES,
                title="СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ",
            ),
            BibliographyEntryNode(number=1, text="Первый источник."),
        ),
        name="keep_sources.docx",
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    heading = doc.paragraphs[texts.index("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ")]
    assert heading.paragraph_format.keep_with_next is True
    blank_after = doc.paragraphs[texts.index("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ") + 1]
    assert blank_after.text == ""
    assert blank_after.paragraph_format.keep_with_next is True


def test_heading_with_figure_keeps_image_and_caption_together(tmp_path):
    images = tmp_path / "images"
    images.mkdir()
    Image.new("RGB", (20, 10), color="white").save(str(images / "figure.png"))

    doc = _render(
        tmp_path,
        (
            HeadingNode(level=HeadingLevel.H2, text="Схема", number="auto"),
            FigureNode(src="figure.png", caption="Схема процесса", id="fig:one"),
        ),
        name="keep_figure.docx",
    )

    texts = [paragraph.text for paragraph in doc.paragraphs]
    heading = doc.paragraphs[texts.index("1.1 Схема")]
    assert heading.paragraph_format.keep_with_next is True

    blank_after = doc.paragraphs[texts.index("1.1 Схема") + 1]
    assert blank_after.text == ""
    assert blank_after.paragraph_format.keep_with_next is True

    image_paragraph = _paragraph_with_drawing(doc)
    assert image_paragraph.paragraph_format.keep_with_next is True

    caption = doc.paragraphs[texts.index("Рисунок 1.1 — Схема процесса")]
    assert caption.paragraph_format.keep_with_next is not True
