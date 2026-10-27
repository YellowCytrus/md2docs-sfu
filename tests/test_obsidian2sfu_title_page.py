"""Regression tests for the title page / ToC page break.

A standalone empty paragraph with a page break used to be left between the title
page and the table of contents.  When the title page filled the sheet, that
paragraph landed on the next page and produced a blank page.  The break now
belongs to the last title line ("Красноярск 2026").
"""

from __future__ import annotations

import importlib.util
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

_REPO_ROOT = Path(__file__).resolve().parents[1]
_MODULE_PATH = _REPO_ROOT / "obsidian2sfu" / "obsidian2sfu.py"
_TITLE_TEMPLATE = _REPO_ROOT / "obsidian2sfu" / "title_reference.docx"


def _load_converter_module():
    spec = importlib.util.spec_from_file_location("obsidian2sfu_under_test", _MODULE_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _page_breaks(paragraph) -> int:
    return sum(
        1 for br in paragraph._p.findall(".//" + qn("w:br")) if br.get(qn("w:type")) == "page"
    )


def _generate_toc_prefixed_document() -> Document:
    doc = Document()
    doc.add_paragraph("city=Красноярск; year=2026")
    doc.add_paragraph("Сгенерированный титульный лист")
    doc.add_page_break()
    doc.add_page_break()
    doc.add_paragraph("СОДЕРЖАНИЕ")
    doc.add_paragraph("Введение")
    return doc


def test_title_page_keeps_single_break_on_the_year_line():
    module = _load_converter_module()

    elements = module.build_title_page(_TITLE_TEMPLATE, {})

    paragraphs = [element for element in elements if element.tag == qn("w:p")]
    breaks = [br for element in elements for br in element.findall(".//" + qn("w:br"))]
    assert len(breaks) == 1
    assert paragraphs[-1].findall(".//" + qn("w:br")) == breaks
    assert "Красноярск" in "".join(t.text or "" for t in paragraphs[-1].findall(".//" + qn("w:t")))


def test_replace_title_page_leaves_no_break_only_paragraph_before_toc():
    module = _load_converter_module()
    doc = _generate_toc_prefixed_document()

    module.replace_title_page(doc, _TITLE_TEMPLATE, {})

    paragraphs = doc.paragraphs
    assert not any(paragraph.text.startswith("Сгенерированный") for paragraph in paragraphs)
    assert not any(paragraph.text.startswith("city=") for paragraph in paragraphs)
    year_index = next(
        index for index, paragraph in enumerate(paragraphs) if "Красноярск" in paragraph.text
    )
    assert paragraphs[year_index + 1].text.strip().upper() == "СОДЕРЖАНИЕ"
    assert _page_breaks(paragraphs[year_index]) == 1
    assert not any(
        _page_breaks(paragraph) and not paragraph.text.strip() for paragraph in paragraphs
    )


def test_replace_title_page_without_toc_falls_back_to_first_break():
    module = _load_converter_module()
    doc = Document()
    doc.add_paragraph("Сгенерированный титульный лист")
    doc.add_page_break()
    doc.add_paragraph("Первый раздел")

    module.replace_title_page(doc, _TITLE_TEMPLATE, {})

    paragraphs = doc.paragraphs
    assert not any(paragraph.text.startswith("Сгенерированный") for paragraph in paragraphs)
    assert paragraphs[-1].text == "Первый раздел"
    assert _page_breaks(paragraphs[-2]) == 1
