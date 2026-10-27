"""Page breaks must never leave a blank page behind.

An empty paragraph that carries nothing but a page break costs a full line.
When the preceding content fills the page exactly, Word pushes that line onto
the next page and applies the break after it, so the reader gets a page with
nothing but the page number. The renderer therefore expresses the break as
``w:pageBreakBefore`` on the paragraph that starts the new page (or appends it
to the previous paragraph when a table comes next) and never leaves a
break-only paragraph in the document.
"""

from __future__ import annotations

from docx import Document as DocxDocument
from docx.oxml.ns import qn

from sfu_converter.domain.ast_nodes import (
    Document,
    HeadingLevel,
    HeadingNode,
    PageBreakNode,
    ParagraphNode,
    TableCell,
    TableNode,
    TableRow,
    TextRun,
)
from sfu_converter.infrastructure.docx_renderer import DocxRenderer
from sfu_converter.registry import get_profile


def _render(tmp_path, blocks, name="breaks.docx"):
    path = tmp_path / name
    DocxRenderer(base_dir=tmp_path).render_to_file(
        Document(blocks=tuple(blocks)),
        get_profile("common"),
        str(path),
    )
    return DocxDocument(str(path))


def _break_only_paragraphs(doc):
    result = []
    for paragraph in doc.paragraphs:
        if not paragraph.text.strip() and any(
            br.get(qn("w:type")) == "page" for br in paragraph._p.iter(qn("w:br"))
        ):
            result.append(paragraph)
    return result


def test_h1_heading_uses_page_break_before_instead_of_break_paragraph(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Конец введения."),)),
            HeadingNode(level=HeadingLevel.H1, text="Позиционирование", number="auto"),
            ParagraphNode(runs=(TextRun("Текст главы."),)),
        ),
    )

    assert _break_only_paragraphs(doc) == []
    heading = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "1 Позиционирование")
    assert heading.paragraph_format.page_break_before is True


def test_explicit_page_break_moves_to_following_paragraph(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Первый блок."),)),
            PageBreakNode(),
            ParagraphNode(runs=(TextRun("Второй блок."),)),
        ),
    )

    assert _break_only_paragraphs(doc) == []
    second = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "Второй блок.")
    assert second.paragraph_format.page_break_before is True


def test_explicit_page_break_before_table_lands_on_caption(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Текст перед таблицей."),)),
            PageBreakNode(),
            TableNode(
                caption="Данные",
                rows=(
                    TableRow(cells=(TableCell("A"), TableCell("B"))),
                    TableRow(cells=(TableCell("1"), TableCell("2"))),
                ),
            ),
        ),
    )

    assert _break_only_paragraphs(doc) == []
    caption = next(paragraph for paragraph in doc.paragraphs if paragraph.text.startswith("Таблица"))
    assert caption.paragraph_format.page_break_before is True


def test_explicit_page_break_before_captionless_table_attaches_to_previous_paragraph(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Текст перед таблицей."),)),
            PageBreakNode(),
            TableNode(
                rows=(
                    TableRow(cells=(TableCell("A"), TableCell("B"))),
                    TableRow(cells=(TableCell("1"), TableCell("2"))),
                ),
            ),
        ),
    )

    assert _break_only_paragraphs(doc) == []
    previous = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "Текст перед таблицей.")
    assert any(br.get(qn("w:type")) == "page" for br in previous._p.iter(qn("w:br")))


def test_trailing_page_break_is_dropped(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Единственный блок."),)),
            PageBreakNode(),
        ),
    )

    assert _break_only_paragraphs(doc) == []
    assert [paragraph.text for paragraph in doc.paragraphs if paragraph.text] == ["Единственный блок."]


def test_consecutive_page_breaks_collapse_to_one(tmp_path):
    doc = _render(
        tmp_path,
        (
            ParagraphNode(runs=(TextRun("Первый блок."),)),
            PageBreakNode(),
            PageBreakNode(),
            ParagraphNode(runs=(TextRun("Второй блок."),)),
        ),
    )

    assert _break_only_paragraphs(doc) == []
    second = next(paragraph for paragraph in doc.paragraphs if paragraph.text == "Второй блок.")
    assert second.paragraph_format.page_break_before is True
