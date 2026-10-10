"""LaTeX math support: inline splitting, V2 parsing, OMML rendering."""

import importlib.util
import shutil
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH

from sfu_converter.domain.ast_nodes import (
    Document,
    FormulaNode,
    ParagraphNode,
    TableCell,
    TableNode,
    TableRow,
    TextRun,
)
from sfu_converter.infrastructure.docx_renderer import DocxRenderer
from sfu_converter.infrastructure.latex_math import has_inline_math, split_inline_math
from sfu_converter.parser.v2_parser import V2Parser
from sfu_converter.registry import get_profile

PANDOC_MISSING = shutil.which("pandoc") is None
requires_pandoc = pytest.mark.skipif(PANDOC_MISSING, reason="pandoc is not installed")

REPO_ROOT = Path(__file__).resolve().parents[1]


def test_split_inline_math_extracts_dollar_regions():
    assert split_inline_math("доход $x^2$ растёт") == [
        ("text", "доход "),
        ("math", "x^2"),
        ("text", " растёт"),
    ]


def test_split_inline_math_supports_paren_delimiters():
    assert split_inline_math(r"\(a + b\)") == [("math", "a + b")]


def test_split_inline_math_without_math_returns_single_text_segment():
    assert split_inline_math("просто текст") == [("text", "просто текст")]


def test_has_inline_math_detects_regions():
    assert has_inline_math("$x$")
    assert not has_inline_math("100 $ на счёт")
    assert not has_inline_math("просто текст")


def test_v2_parser_reads_formula_format_attribute():
    source = "\n".join(
        [
            "[DOC syntax=2 profile=common language=ru]",
            "[FORMULA id=eq:main number=auto format=latex]",
            "E = mc^2",
            "[FORMULA_END]",
        ]
    )
    result = V2Parser().parse(source)
    formula = next(block for block in result.document.blocks if isinstance(block, FormulaNode))
    assert formula.format == "latex"
    assert formula.content == "E = mc^2"


@requires_pandoc
def test_renderer_renders_latex_formula_as_equation_with_right_number(tmp_path):
    renderer = DocxRenderer(base_dir=tmp_path)
    ast = Document(blocks=(FormulaNode(content=r"E = mc^2", format="latex"),))
    output_path = tmp_path / "latex_formula.docx"

    renderer.render_to_file(ast, get_profile("common"), str(output_path))

    doc = DocxDocument(str(output_path))
    assert "<m:oMath" in doc.element.body.xml
    formula_para = next(p for p in doc.paragraphs if p.text == "\t(1)")
    assert formula_para.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert 'w:val="right"' in formula_para._p.xml


@requires_pandoc
def test_renderer_renders_inline_math_as_equation(tmp_path):
    renderer = DocxRenderer(base_dir=tmp_path)
    ast = Document(
        blocks=(ParagraphNode(runs=(TextRun(text="значение $x^2$ растёт"),)),)
    )
    output_path = tmp_path / "inline_math.docx"

    renderer.render_to_file(ast, get_profile("common"), str(output_path))

    doc = DocxDocument(str(output_path))
    assert "<m:oMath" in doc.element.body.xml
    combined = "".join(p.text for p in doc.paragraphs)
    assert "значение" in combined and "растёт" in combined
    assert "$" not in combined


@requires_pandoc
def test_renderer_renders_inline_math_inside_table_cell(tmp_path):
    renderer = DocxRenderer(base_dir=tmp_path)
    table = TableNode(
        rows=(
            TableRow(cells=(TableCell(text="Показатель"), TableCell(text="Сумма"))),
            TableRow(cells=(TableCell(text="Доход"), TableCell(text="$30\\,000$"))),
        )
    )
    ast = Document(blocks=(table,))
    output_path = tmp_path / "table_math.docx"

    renderer.render_to_file(ast, get_profile("common"), str(output_path))

    doc = DocxDocument(str(output_path))
    assert "<m:oMath" in doc.element.body.xml
    assert "$" not in doc.element.body.xml


def _load_obsidian2sfu():
    spec = importlib.util.spec_from_file_location(
        "obsidian2sfu_module", REPO_ROOT / "obsidian2sfu" / "obsidian2sfu.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _markdown_converter(module):
    config = {
        "meta": {"profile": "lab_practical_project_reports", "language": "ru"},
        "meta_items": [],
        "sources": [],
        "kv_first_cells": [],
        "kv_captions": [],
        "ref_fixes": [],
        "promote_to_h1": [],
        "vault": "/tmp",
    }
    return module.MarkdownConverter(config)


def test_markdown_display_math_becomes_latex_formula_blocks():
    module = _load_obsidian2sfu()
    converter = _markdown_converter(module)
    text = "Текст\n\n$$E = mc^2$$\n\nдве:\n\n$$a + b$$ $$c - d$$\n\nблок:\n\n$$\nFV = PV\n$$\n"

    txt = converter.convert(text, Path("/tmp"))

    assert txt.count("[FORMULA format=latex number=auto]") == 4
    assert txt.count("[FORMULA_END]") == 4
    assert "E = mc^2" in txt
    assert "FV = PV" in txt


def test_markdown_inline_math_is_left_intact():
    module = _load_obsidian2sfu()
    converter = _markdown_converter(module)
    txt = converter.convert("Доход равен $8000 - 2500 = 5500$ руб.\n", Path("/tmp"))
    assert "$8000 - 2500 = 5500$" in txt
